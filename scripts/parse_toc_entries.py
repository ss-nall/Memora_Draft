import sys
import re
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('docs/Memora_Report_Rewritten-1.docx')

# 1. Figures
figures_list = []
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("Figure ") and ("...." in txt or "\t" in txt):
        # Extract figure number, title, page
        m = re.match(r'^(Figure\s+[\d\.]+):\s*(.*?)\.{4,}\s*(\d+)$', txt)
        if m:
            figures_list.append((m.group(1).replace("Figure ", "").strip(), m.group(2).strip(), m.group(3).strip()))
        else:
            # fallback
            parts = txt.split(":")
            fnum = parts[0].replace("Figure ", "").strip()
            rest = parts[1]
            dots_idx = rest.find("...")
            title = rest[:dots_idx].strip() if dots_idx != -1 else rest.strip()
            page = re.findall(r'\d+$', rest)
            page_str = page[0] if page else ""
            figures_list.append((fnum, title, page_str))

print(f"Total Figures extracted: {len(figures_list)}")
for f in figures_list[:5]:
    print("  Fig:", f)

# 2. Tables
tables_list = []
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("Table ") and ("...." in txt or "\t" in txt):
        m = re.match(r'^(Table\s+[\d\.]+):\s*(.*?)\.{4,}\s*(\d+)$', txt)
        if m:
            tables_list.append((m.group(1).replace("Table ", "").strip(), m.group(2).strip(), m.group(3).strip()))
        else:
            parts = txt.split(":")
            tnum = parts[0].replace("Table ", "").strip()
            rest = parts[1]
            dots_idx = rest.find("...")
            title = rest[:dots_idx].strip() if dots_idx != -1 else rest.strip()
            page = re.findall(r'\d+$', rest)
            page_str = page[0] if page else ""
            tables_list.append((tnum, title, page_str))

print(f"\nTotal Tables extracted: {len(tables_list)}")
for t in tables_list[:5]:
    print("  Tbl:", t)

# 3. TOC from SDT
sdt = doc._body._element.findall(docx.oxml.ns.qn('w:sdt'))[0]
sdtContent = sdt.find(docx.oxml.ns.qn('w:sdtContent'))
toc_paras = [p for p in sdtContent.findall(docx.oxml.ns.qn('w:p'))]
print(f"\nTotal TOC paras in SDT: {len(toc_paras)}")
toc_entries = []
for p in toc_paras:
    txt = ''.join(p.itertext()).strip()
    if not txt or txt == "TABLE OF CONTENTS":
        continue
    # Typically formatted as: Title \t Page or Sl.No \t Title \t Page
    # Let's inspect text
    toc_entries.append(txt)

print(f"Total TOC entries: {len(toc_entries)}")
for e in toc_entries[:10]:
    print("  TOC:", repr(e))
