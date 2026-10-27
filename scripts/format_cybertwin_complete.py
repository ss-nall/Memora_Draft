import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client
import pymupdf

sys.stdout.reconfigure(encoding='utf-8')

print("=== BUILDING CYBERTWIN FORMATTED MEMORA REPORT ===")

def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_clean_three_line_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="10" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="10" w:space="0" w:color="000000"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="none"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def add_header_bottom_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def set_table_borders(table, border_color="C0C0C0"):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def clear_all_headers(doc):
    for sec in doc.sections:
        sec.different_first_page_header_footer = False
        hdr = sec.header
        hdr.is_linked_to_previous = False
        for p in list(hdr.paragraphs):
            pPr = p._p.pPr
            p._p.clear()
            if pPr is not None:
                p._p.append(pPr)
        first_hdr = sec.first_page_header
        first_hdr.is_linked_to_previous = False
        for p in list(first_hdr.paragraphs):
            pPr = p._p.pPr
            p._p.clear()
            if pPr is not None:
                p._p.append(pPr)

# Load original docx
doc = docx.Document('docs/Memora_Report_Rewritten-1.docx')

# 1. Clear all headers completely
clear_all_headers(doc)

# 2. Configure margins matching CyberTwin
for sec in doc.sections:
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.17)
    sec.right_margin = Inches(0.95)
    sec.header_distance = Inches(0.5)
    sec.footer_distance = Inches(0.5)

# 3. Configure styles to Times New Roman
for s_name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4']:
    if s_name in doc.styles:
        st = doc.styles[s_name]
        st.font.name = 'Times New Roman'
        if s_name == 'Normal':
            st.font.size = Pt(12)
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.space_after = Pt(5)
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# 4. Extract figures and tables definitions before replacing dotted paragraphs
figures_data = []
tables_data = []

fig_indices_to_remove = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("Figure ") and ("...." in txt or "\t" in txt):
        fig_indices_to_remove.append(i)
        parts = txt.split(":")
        fnum = parts[0].replace("Figure ", "").strip()
        rest = parts[1]
        dots_idx = rest.find("...")
        title = rest[:dots_idx].strip() if dots_idx != -1 else rest.strip()
        page = re.findall(r'\d+$', rest)
        page_str = page[0] if page else ""
        figures_data.append((fnum, title, page_str))

tbl_indices_to_remove = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("Table ") and ("...." in txt or "\t" in txt):
        tbl_indices_to_remove.append(i)
        parts = txt.split(":")
        tnum = parts[0].replace("Table ", "").strip()
        rest = parts[1]
        dots_idx = rest.find("...")
        title = rest[:dots_idx].strip() if dots_idx != -1 else rest.strip()
        page = re.findall(r'\d+$', rest)
        page_str = page[0] if page else ""
        tables_data.append((tnum, title, page_str))

print(f"Extracted {len(figures_data)} figures and {len(tables_data)} tables from lists.")

# Clear the dotted text paragraphs
for idx in fig_indices_to_remove:
    doc.paragraphs[idx].text = ""
for idx in tbl_indices_to_remove:
    doc.paragraphs[idx].text = ""

# 5. Format all headings and content
prelim_titles = [
    "DECLARATION", "CERTIFICATE", "ACKNOWLEDGEMENTS", "TABLE OF CONTENTS",
    "LIST OF FIGURES", "LIST OF TABLES", "LIST OF ABBREVIATIONS",
    "SYMBOLS AND NOTATIONS", "ABSTRACT"
]

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue

    # Preliminary Headings: Centered, 15pt Bold
    if txt in prelim_titles:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(16)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(15)
            r.bold = True
        continue

    # CHAPTER X: Centered, 16pt Bold
    if txt.startswith("CHAPTER "):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(16)
            r.bold = True
        continue

    # Main Chapter Title (Heading 1: Left, 14pt Bold)
    if p.style.name == 'Heading 1' and txt not in prelim_titles:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.bold = True
        continue

    # Section Heading (Heading 2: Left, 14pt Bold)
    if p.style.name == 'Heading 2':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.bold = True
        continue

    # Subsection Heading (Heading 3: Left, 12pt Bold)
    if p.style.name == 'Heading 3':
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.bold = True
        continue

    # Level 4 Headings (Algorithms / RQ questions)
    if p.style.name == 'Heading 4' or txt.startswith("Algorithm ") or (txt.startswith("RQ") and ":" in txt):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.bold = True
        continue

    # Figure and Table Captions
    if txt.startswith("Figure "):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
        continue

    if txt.startswith("Table ") and (":" in txt or "\t" in txt) and len(txt) < 130:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.bold = True
        continue

    # References entries
    if re.match(r'^\[\d+\]', txt):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
        continue

    # Regular body paragraphs
    if p.style.name == 'Normal' and i > 25:
        if p._p.xpath('.//a:blip'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.space_before = Pt(0)
            for r in p.runs:
                if r.font.name != 'Consolas':
                    r.font.name = 'Times New Roman'
                    if not r.font.size:
                        r.font.size = Pt(12)

# Format all data tables
for t_idx, tbl in enumerate(doc.tables):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    if t_idx in [0, 1, 2]:
        for row in tbl.rows:
            for c in row.cells:
                set_cell_margins(c, top=40, bottom=40, left=60, right=60)
                for cp in c.paragraphs:
                    for cr in cp.runs:
                        cr.font.name = 'Times New Roman'
        continue

    set_table_borders(tbl, "C0C0C0")
    for r_idx, row in enumerate(tbl.rows):
        is_header = (r_idx == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c in row.cells:
            set_cell_margins(c, top=45, bottom=45, left=75, right=75)
            if is_header:
                add_header_bottom_border(c)
            
            for cp in c.paragraphs:
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(1)
                cp.paragraph_format.line_spacing = 1.05
                for cr in cp.runs:
                    cr.font.name = 'Times New Roman'
                    if is_header:
                        cr.bold = True
                        cr.font.size = Pt(10)
                    else:
                        if not cr.font.size or cr.font.size.pt > 10.5:
                            cr.font.size = Pt(9.5)

# Save intermediate docx
temp_docx = os.path.abspath('docs/temp_pass1.docx')
doc.save(temp_docx)
print("Saved intermediate docx.")

# Export to PDF to scan actual page positions
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
temp_pdf = os.path.abspath('docs/temp_pass1.pdf')
wdoc = word.Documents.Open(temp_docx)
wdoc.SaveAs(temp_pdf, FileFormat=17)
wdoc.Close()
word.Quit()
print("Saved intermediate PDF for page scanning.")

# Scan PDF for exact page numbers
pdf = pymupdf.open(temp_pdf)
print(f"Scanned intermediate PDF pages: {len(pdf)}")

page_for_item = {}
for pno in range(len(pdf)):
    page = pdf[pno]
    txt = page.get_text()
    
    # Check printed footer page number
    blocks = page.get_text('dict')['blocks']
    footer_num = ""
    for b in blocks:
        if 'lines' in b:
            for l in b['lines']:
                for s in l['spans']:
                    if s['bbox'][1] > 780 and s['text'].strip():
                        footer_num = s['text'].strip()
    
    for line in txt.split('\n'):
        line_clean = line.strip()
        if not line_clean:
            continue
        # Check headings
        for prefix in ["1. INTRODUCTION", "1.1 Background", "1.2 Motivations", "1.3 Scope of the Project",
                       "2. PROJECT DESCRIPTION AND GOALS", "2.1 Literature Review", "2.2 Research Gap", "2.3 Objectives",
                       "2.4 Problem Statement", "2.5 Project Plan",
                       "3. TECHNICAL SPECIFICATION", "3.1 Requirements", "3.2 Feasibility Study", "3.3 System Specification",
                       "4. DESIGN APPROACH AND DETAILS", "4.1 System Architecture", "4.2 Design",
                       "5. METHODOLOGY AND TESTING", "5.1 Module Description", "5.2 Experimental Methodology", "5.3 Testing",
                       "6. PROJECT DEMONSTRATION", "6.1 Web Dashboard", "6.2 Zero-Install Guest Onboarding",
                       "6.3 Automated Photo Matching", "6.4 Automated Memory Reel Synthesis", "6.5 Cross-Platform Mobile Deployment",
                       "7. RESULT AND DISCUSSION", "7.1 Experimental Results", "7.2 Discussion", "7.3 Cost Analysis",
                       "8. CONCLUSION", "8.1 Summary of Contributions", "8.2 Societal and Ethical Impact",
                       "8.3 Limitations and Future Research Directions", "9. REFERENCES", "APPENDIX A", "ABSTRACT"]:
            if prefix in line_clean and prefix not in page_for_item and pno >= 10:
                page_for_item[prefix] = footer_num if footer_num else str(pno + 1)
        
        # Check figures
        m_f = re.match(r'^(Figure\s+\d+\.\d+)', line_clean)
        if m_f and pno >= 12:
            fkey = m_f.group(1).replace("Figure ", "").strip()
            if fkey not in page_for_item:
                page_for_item[fkey] = footer_num if footer_num else str(pno + 1)

        # Check tables
        m_t = re.match(r'^(Table\s+\d+\.\d+)', line_clean)
        if m_t and pno >= 12:
            tkey = m_t.group(1).replace("Table ", "").strip()
            if tkey not in page_for_item:
                page_for_item[tkey] = footer_num if footer_num else str(pno + 1)

print(f"Scanned {len(page_for_item)} item page locations.")

# Clean up temporary pass 1 files
pdf.close()
try:
    os.remove(temp_pdf)
    os.remove(temp_docx)
except Exception:
    pass

print("Pass 1 complete. Now building CyberTwin TOC and Lists tables...")
